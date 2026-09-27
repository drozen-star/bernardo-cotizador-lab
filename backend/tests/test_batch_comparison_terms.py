"""Lab L5f (fase 3): régimen y flete por proveedor en el comparativo. Caso numérico de referencia."""

# ruff: noqa: F811 - fixtures importados y recibidos como parámetro

from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.features.batch_comparison import freight
from app.features.batch_comparison import regime
from app.features.batch_comparison import strategies
from app.features.batch_comparison.fiscal import format_money
from app.features.batch_comparison.regime import SupplierTerms
from app.features.whatsapp.model import WhatsappConversation
from tests.test_batch_comparison import _cost
from tests.test_batch_comparison import _rfq
from tests.test_batch_comparison_api import HEADERS
from tests.test_batch_comparison_api import world  # noqa: F401 - fixture

A, B, C = "id:1", "id:2", "id:3"
NAMES = {A: "Corralón A", B: "Corralón B", C: "Corralón C"}

#: Batch 1: (rfq_id, nombre, cantidad)
ITEMS = [(1, "Ladrillo", 1200), (2, "Cemento", 60), (3, "Cal", 40), (4, "Arena", 6), (5, "Hierro", 150)]
PRICES_A = {1: "1180", 2: "11950", 3: "6420", 4: "38500", 5: "9870"}
PRICES_B = {3: "6300", 4: "37000"}

TERMS_A = SupplierTerms(billing_regime="facturado", freight_cost=Decimal("45000"), freight_basis="pedido")
TERMS_B = SupplierTerms(billing_regime="facturado", freight_cost=Decimal("40000"), freight_basis="pedido")


def _items():
    return [_rfq(rfq_id, name, quantity) for rfq_id, name, quantity in ITEMS]


def _costs(items, supplier, prices, terms):
    return [
        _cost(prices[rfq.id], rfq, supplier, iva_included=False, freight_included=None, terms=terms)
        for rfq in items if rfq.id in prices
    ]


# ------------------------------------------------------------ referencia
def test_reference_case_everything_to_a_wins_over_the_split():
    items = _items()
    costs = _costs(items, A, PRICES_A, TERMS_A) + _costs(items, B, PRICES_B, TERMS_B)

    comparison = strategies.compare(items, costs, NAMES, {A: TERMS_A, B: TERMS_B})
    lowest = comparison.lowest_cost

    assert lowest.key == "menor_costo_total" and lowest.title == "Menor costo total"
    assert lowest.supplier_names == ["Corralón A"]
    assert lowest.total_costo_real == Decimal("4146300")  # 4.101.300 de ítems + 45.000 de flete
    assert lowest.freight_by_supplier == {"Corralón A": Decimal("45000")}
    assert lowest.freight_desembolso_by_supplier == {"Corralón A": Decimal("54450")}  # facturado: flete con IVA
    assert lowest.total_desembolso == Decimal("4101300") * Decimal("1.21") + Decimal("54450")
    assert "El flete suma $ 45.000,00 (Corralón A $ 45.000,00)." in lowest.porque

    # La división A + B (cal y arena a B) da 4.172.500 y no gana.
    split = strategies._subset_total([rfq for rfq in items], strategies.build_matrix(costs), {A, B}, NAMES, {A: TERMS_A, B: TERMS_B})
    assert split[0] == Decimal("4172500")

    fewer = comparison.fewer_suppliers
    assert fewer.supplier_names == ["Corralón A"] and fewer.total_costo_real == Decimal("4146300")
    assert comparison.difference_amount == Decimal("0")
    assert not any(mark.startswith("Corralón") for mark in lowest.marks)  # flete por pedido, sin umbral: sin marca


def test_reference_case_threshold_makes_freight_free():
    items = _items()
    terms_a = SupplierTerms(billing_regime="facturado", freight_cost=Decimal("30000"), freight_free_over=Decimal("500000"))
    costs = _costs(items, A, PRICES_A, terms_a) + _costs(items, B, PRICES_B, TERMS_B)

    lowest = strategies.compare(items, costs, NAMES, {A: terms_a, B: TERMS_B}).lowest_cost

    assert lowest.supplier_names == ["Corralón A"]
    assert lowest.total_costo_real == Decimal("4101300")
    assert lowest.freight_by_supplier == {"Corralón A": Decimal("0")}
    assert "Corralón A: flete sin cargo: pedido de $ 4.101.300,00 supera $ 500.000,00" in lowest.marks
    assert "flete" not in lowest.porque.split("El plazo")[0]  # sin flete que pese, no se nombra


def test_threshold_not_reached_charges_freight_with_mark_and_evaluates_only_what_is_assigned():
    items = _items()
    terms_a = SupplierTerms(billing_regime="facturado", freight_cost=Decimal("30000"), freight_free_over=Decimal("5000000"))
    only_cal = [rfq for rfq in items if rfq.id == 3]
    costs = _costs(only_cal, A, PRICES_A, terms_a)

    amount, marks = freight.freight_for(A, costs, terms_a)

    assert amount == Decimal("30000")
    assert marks == ["flete $ 30.000,00: pedido de $ 256.800,00 no llega a $ 5.000.000,00"]

    no_cost = SupplierTerms(billing_regime="facturado", freight_free_over=Decimal("5000000"))
    assert freight.freight_for(A, costs, no_cost) == (Decimal("0"), ["flete a cotizar por debajo de $ 5.000.000,00"])


def test_freight_per_trip_assumes_one_trip_and_included_is_zero():
    cal = [_cost("6420", _rfq(3, "Cal", 40), A, iva_included=False, terms=TERMS_A)]

    trip = SupplierTerms(billing_regime="facturado", freight_cost=Decimal("30000"), freight_basis="viaje")
    assert freight.freight_for(A, cal, trip) == (Decimal("30000"), ["flete $ 30.000,00 por viaje: 1 viaje supuesto"])

    included = SupplierTerms(billing_regime="facturado", freight_cost=Decimal("0"), freight_basis="pedido")
    assert freight.freight_for(A, cal, included) == (Decimal("0"), [])
    assert freight.freight_for(A, cal, None) == (Decimal("0"), [])

    # Sin dato de flete: 0 con la marca del ítem.
    unknown = [_cost("6420", _rfq(3, "Cal", 40), A, iva_included=False, freight_included=False, terms=SupplierTerms(billing_regime="facturado"))]
    assert freight.freight_for(A, unknown, SupplierTerms(billing_regime="facturado")) == (Decimal("0"), ["flete a cotizar"])

    assert freight.freight_disbursement(Decimal("45000"), TERMS_A) == Decimal("54450")
    assert freight.freight_disbursement(Decimal("45000"), SupplierTerms(billing_regime="efectivo")) == Decimal("45000")
    assert freight.freight_disbursement(Decimal("0"), TERMS_A) == Decimal("0")


# ---------------------------------------------------------------- régimen
def test_cash_regime_price_is_the_disbursement_and_the_real_cost():
    cal = _rfq(3, "Cal", 40)
    cash = SupplierTerms(billing_regime="efectivo", freight_cost=Decimal("40000"))

    cost = _cost("7500", cal, C, iva_included=None, freight_included=None, terms=cash)

    assert cost.desembolso_total == Decimal("300000") and cost.costo_real_total == Decimal("300000")
    assert cost.billing_regime == "efectivo"
    assert regime.MARK_CASH in cost.marks
    assert regime.MARK_IVA_UNCONFIRMED not in cost.marks  # con efectivo el IVA no se pide ni se marca
    assert "flete sin confirmar" not in cost.marks  # el flete lo dice la conversación


def test_partial_regime_with_and_without_percentage():
    rfq = _rfq(9, "Bloque", 1)
    half = SupplierTerms(billing_regime="parcial", documented_pct=Decimal("50"))

    cost = _cost("121000", rfq, A, iva_included=True, terms=half)
    assert cost.desembolso_total == Decimal("121000")
    assert cost.costo_real_total == Decimal("110500")  # 121.000 − (60.500 / 1,21) × 0,21
    assert "factura el 50%" in cost.marks

    unknown = _cost("121000", rfq, A, iva_included=True, terms=SupplierTerms(billing_regime="parcial"))
    assert unknown.costo_real_total == Decimal("121000")
    assert regime.MARK_PARTIAL_UNKNOWN in unknown.marks


def test_missing_regime_is_assumed_invoiced_with_mark():
    cost = _cost("1210", terms=None)
    assert cost.neto_unit == Decimal("1000") and regime.MARK_REGIME_UNCONFIRMED in cost.marks

    with_conversation = _cost("1210", terms=SupplierTerms())
    assert regime.MARK_REGIME_UNCONFIRMED in with_conversation.marks
    assert regime.MARK_REGIME_UNCONFIRMED not in _cost("1210").marks


def test_mixed_regimes_are_marked_and_consolidating_saves_freight():
    items = _items()
    cash_c = SupplierTerms(billing_regime="efectivo", freight_cost=Decimal("40000"), freight_basis="pedido")
    costs = _costs(items, A, PRICES_A, TERMS_A) + [_cost("5000", items[2], C, iva_included=None, terms=cash_c)]

    comparison = strategies.compare(items, costs, NAMES, {A: TERMS_A, C: cash_c})
    lowest, fewer = comparison.lowest_cost, comparison.fewer_suppliers

    # Cal a C en efectivo: 200.000 + 40.000 de flete < 256.800 de A.
    assert lowest.supplier_names == ["Corralón A", "Corralón C"]
    assert lowest.total_costo_real == Decimal("4146300") - Decimal("256800") + Decimal("240000")
    assert strategies.MARK_MIXED_REGIMES in lowest.marks
    assert lowest.freight_desembolso_by_supplier["Corralón C"] == Decimal("40000")  # efectivo: flete tal cual

    assert fewer.supplier_names == ["Corralón A"] and strategies.MARK_MIXED_REGIMES not in fewer.marks
    assert "Consolidar en Corralón A ahorra $ 40.000,00 de flete neto." in fewer.porque
    assert format_money(comparison.difference_amount) == "$ 16.800,00"


# ------------------------------------------------------------------ API
def test_comparison_api_exposes_terms_regime_column_and_freight_rows(client, db_session, world):
    conversation = db_session.query(WhatsappConversation).one()  # la de Corralón Norte
    conversation.billing_regime = "efectivo"
    conversation.freight_cost = Decimal("30000")
    conversation.freight_basis = "viaje"
    db_session.commit()

    body = client.get(f"/rfq-batches/{world.batch.id}/comparison", headers=HEADERS).json()
    norte = body["suppliers"][0]
    assert norte["name"] == "Corralón Norte" and norte["billing_regime"] == "efectivo"
    assert norte["freight"] == {"included": False, "cost": "30000.00", "basis": "viaje", "free_over": None}
    ladrillo_norte = body["items"][0]["quotes"][f"id:{world.norte.id}"]
    assert ladrillo_norte["billing_regime"] == "efectivo" and ladrillo_norte["costo_real_total"] == "1452000.00"  # 1.210 × 1.200 tal cual
    assert "en efectivo, sin factura" in ladrillo_norte["marks"]

    lowest = body["strategies"][0]
    assert lowest["freight_by_supplier"].get("Corralón Norte") == "30000.00"
    assert any("1 viaje supuesto" in mark for mark in lowest["marks"])
    assert "mezcla facturado y efectivo: comparar con cuidado" in lowest["marks"]

    workbook = load_workbook(BytesIO(client.get(f"/rfq-batches/{world.batch.id}/comparison.xlsx", headers=HEADERS).content))
    matrix = workbook["Matriz"]
    assert matrix.cell(row=3, column=5).value == "efectivo"
    assert matrix.cell(row=3, column=10).value == "$ 30.000,00 por viaje"
    texts = [str(c.value) for row in workbook["Estrategias"].iter_rows() for c in row if c.value]
    assert "Flete Corralón Norte" in texts
    assumptions = [str(c.value) for row in workbook["Supuestos"].iter_rows() for c in row if c.value]
    assert any("Régimen por proveedor" in t for t in assumptions) and any("una vez por proveedor" in t for t in assumptions)
    assert not any("se asumen facturadas" in t for t in assumptions)
