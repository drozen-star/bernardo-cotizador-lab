"""Lab L5d: loader, endpoints del comparativo, Excel y script demo (SQLite del conftest)."""

from datetime import date
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.mixins import utcnow
from app.features.auth.model import User
from app.features.batch_comparison import excel
from app.features.batch_comparison.loader import load_batch
from app.features.quote.model import SupplierQuote
from app.features.rfq.batch_service import create_batch_from_rows
from app.features.rfq.intake import parse_materials_xlsx
from app.features.supplier.model import Supplier
from app.features.whatsapp.model import WhatsappConversation
from app.features.whatsapp.settings import whatsapp_settings
from scripts import lab_demo_comparison as demo
from tests.test_rfq_batch import OK_FILE
from tests.test_rfq_batch import _create_supplier
from tests.test_rfq_batch import _register

ADMIN = "admin-token-comparativo"
HEADERS = {"X-Bernardo-Lab-Admin": ADMIN}


class World:
    def __init__(self, user, batch, norte, sur):
        self.user, self.batch, self.norte, self.sur = user, batch, norte, sur
        self.rfqs = sorted(batch.rfqs, key=lambda rfq: rfq.id)
        self.ladrillo, self.cemento = self.rfqs[0], self.rfqs[1]


def _quote(rfq, supplier, price, *, submitted_at=None, **overrides):
    fields = dict(
        rfq_id=rfq.id, supplier_id=supplier.id, supplier_name=supplier.name,
        unit_price=Decimal(price) if price is not None else None, currency="ARS", unit=rfq.unit,
        source="manual", iva_included=True, freight_included=True, lead_time=3, payment_terms="contado",
        validity_date=date(2026, 10, 15), submitted_at=submitted_at or utcnow(),
    )
    fields.update(overrides)

    return SupplierQuote(**fields)


@pytest.fixture
def world(client, db_session, monkeypatch):
    monkeypatch.setattr(whatsapp_settings, "LAB_ADMIN_TOKEN", ADMIN)
    data = _register(client, "compras@constructora.example.com")
    user = db_session.get(User, data["user"]["id"])
    norte = db_session.get(Supplier, _create_supplier(client, data["access_token"], "Corralón Norte", "n@x.example.com")["id"])
    sur = db_session.get(Supplier, _create_supplier(client, data["access_token"], "Materiales del Sur", "s@x.example.com")["id"])

    batch = create_batch_from_rows(
        db=db_session, user=user, rows=parse_materials_xlsx(OK_FILE), name="Obra Palermo - mampostería",
        site_name="Edificio Palermo Soho", site_address="Gorriti 4800, CABA", delivery_expectation=date(2026, 10, 15),
    )
    world = World(user, batch, norte, sur)
    old = utcnow() - timedelta(days=1)

    db_session.add_all([
        _quote(world.ladrillo, norte, "999", submitted_at=old),  # vieja: la pisa la siguiente
        _quote(world.ladrillo, norte, "1210"),
        _quote(world.cemento, norte, "12100", iva_included=False),
        _quote(world.ladrillo, sur, "900", iva_included=False, freight_included=False, lead_time=5),
        _quote(world.cemento, sur, "20", currency="USD"),
        WhatsappConversation(rfq_batch_id=batch.id, supplier_id=norte.id, status="complete"),
    ])
    db_session.commit()

    return world


def _get(client, batch_id, suffix="", headers=HEADERS, **params):
    return client.get(f"/rfq-batches/{batch_id}/comparison{suffix}", headers=headers, params=params)


# ------------------------------------------------------------------- loader
def test_loader_keeps_latest_quote_per_supplier_and_conversation_status(db_session, world):
    loaded = load_batch(db_session, world.batch.id)

    assert [rfq.id for rfq in loaded.items] == [rfq.id for rfq in world.rfqs]
    assert [s.name for s in loaded.suppliers] == ["Corralón Norte", "Materiales del Sur"]
    assert loaded.suppliers[0].conversation_status == "complete"
    assert loaded.suppliers[1].conversation_status is None
    assert loaded.quote_for(world.ladrillo.id, f"id:{world.norte.id}").unit_price == Decimal("1210")
    assert loaded.quote_for(world.rfqs[2].id, f"id:{world.norte.id}") is None
    assert load_batch(db_session, 9999) is None


# ------------------------------------------------------------------- router
def test_comparison_requires_admin_token(client, world):
    assert _get(client, world.batch.id, headers={}).status_code == 401
    assert _get(client, world.batch.id, headers={"X-Bernardo-Lab-Admin": "no"}).status_code == 401
    assert _get(client, world.batch.id, ".xlsx", headers={}).status_code == 401


def test_comparison_unknown_batch_is_404(client, world):
    assert _get(client, 9999).status_code == 404
    assert _get(client, 9999, ".xlsx").status_code == 404


@pytest.mark.parametrize("raw", ["abc", "1:", "1:x", "1:150"])
def test_comparison_malformed_alicuotas_is_422(client, world, raw):
    assert _get(client, world.batch.id, alicuotas=raw).status_code == 422


def test_comparison_alicuota_for_foreign_item_is_422(client, world):
    response = _get(client, world.batch.id, alicuotas="9999:10.5")

    assert response.status_code == 422
    assert "9999" in response.json()["detail"]


def test_comparison_json(client, world):
    response = _get(client, world.batch.id)

    assert response.status_code == 200
    body = response.json()
    norte, sur = f"id:{world.norte.id}", f"id:{world.sur.id}"

    assert body["batch"]["name"] == "Obra Palermo - mampostería"
    assert body["currency"] == "ARS"
    assert [s["name"] for s in body["suppliers"]] == ["Corralón Norte", "Materiales del Sur"]
    assert body["suppliers"][0]["conversation_status"] == "complete"
    assert len(body["items"]) == 5

    ladrillo = body["items"][0]
    assert ladrillo["alicuota"] == "21"
    assert ladrillo["quotes"][norte]["neto_unit"] == "1000.00"
    assert ladrillo["quotes"][norte]["costo_real_total"] == "1200000.00"  # 1000 × 1200
    assert ladrillo["quotes"][sur]["costo_real_total"] == "1080000.00"
    # L5f: sin conversación no hay régimen: se asume facturado con marca.
    assert ladrillo["quotes"][sur]["marks"] == ["régimen sin confirmar, se asume facturado", "flete a cotizar"]
    assert ladrillo["quotes"][sur]["billing_regime"] is None
    assert body["suppliers"][1]["freight"] is None and body["suppliers"][0]["billing_regime"] is None
    assert ladrillo["best_supplier_key"] == sur

    cemento = body["items"][1]
    assert cemento["quotes"][sur]["comparable"] is False
    assert cemento["quotes"][sur]["marks"] == ["moneda USD, no comparada"]
    assert cemento["best_supplier_key"] == norte

    assert [s["key"] for s in body["strategies"]] == ["menor_costo_total", "menos_proveedores"]  # L5f
    lowest, fewer = body["strategies"]
    assert lowest["total_costo_real"] == "1806000.00"  # 1.080.000 + 726.000
    assert lowest["supplier_count"] == 2
    assert fewer["supplier_names"] == ["Corralón Norte"] and fewer["total_costo_real"] == "1926000.00"
    assert body["difference"] == {"amount": "120000.00", "pct": "6.6"}
    assert "pagás $ 120.000,00 más (6,6%)" in fewer["porque"]
    assert all(isinstance(v, str) for v in (lowest["total_costo_real"], lowest["total_desembolso"]))


def test_comparison_alicuota_override_changes_the_net(client, world):
    response = _get(client, world.batch.id, alicuotas=f"{world.ladrillo.id}:10.5")

    assert response.status_code == 200
    body = response.json()
    assert body["alicuotas"][str(world.ladrillo.id)] == "10.5"
    assert body["items"][0]["quotes"][f"id:{world.norte.id}"]["neto_unit"] == "1095.02"
    assert body["items"][1]["alicuota"] == "21"


def test_comparison_xlsx(client, world):
    response = _get(client, world.batch.id, ".xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(excel.XLSX_MEDIA_TYPE)
    disposition = response.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="comparativo-obra-palermo-mamposteria-')
    assert disposition.endswith('.xlsx"')

    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Matriz", "Estrategias", "Supuestos"]

    matrix = workbook["Matriz"]
    assert matrix.cell(row=1, column=4).value == "Corralón Norte (WhatsApp: complete)"
    assert matrix.cell(row=2, column=1).value == "Ítem"
    assert matrix.cell(row=3, column=1).value == world.ladrillo.item_name
    price = matrix.cell(row=3, column=4)
    assert isinstance(price.value, (int, float)) and price.value == 1210
    assert price.number_format == "#,##0.00"
    assert matrix.cell(row=2, column=5).value == "Régimen" and matrix.cell(row=3, column=5).value == "sin confirmar"  # L5f
    total_norte = matrix.cell(row=3, column=8)
    total_sur = matrix.cell(row=3, column=19)
    assert total_norte.value == 1200000 and total_sur.value == 1080000
    assert total_sur.fill.fgColor.rgb.endswith("C6EFCE") and total_norte.fill.fill_type is None

    texts = [str(cell.value) for row in workbook["Supuestos"].iter_rows() for cell in row if cell.value]
    assert any("Régimen por proveedor" in text for text in texts)  # L5f: reemplaza "se asumen facturadas"
    assert any("Buenos Aires" in text for text in texts)
    assert any(text.startswith("Alícuota IVA:") for text in texts)
    assert any("ARS" in text for text in texts)

    strategy_texts = [str(cell.value) for row in workbook["Estrategias"].iter_rows() for cell in row if cell.value]
    assert "Menor costo total" in strategy_texts and "Menos proveedores" in strategy_texts  # L5f
    assert "Diferencia entre estrategias" in strategy_texts


def test_filename_slug_has_no_odd_characters():
    assert excel.slugify("Obra Palermo - mampostería / hierro ñ") == "obra-palermo-mamposteria-hierro-n"
    assert excel.slugify("¿?") == "pedido"


# ------------------------------------------------------------------- script
def test_demo_script_refuses_non_sqlite(capsys, monkeypatch):
    with pytest.raises(SystemExit) as exc:
        demo.refuse_unless_sqlite("postgresql+psycopg://user@host/db")

    assert exc.value.code == 1
    assert demo.REFUSAL in capsys.readouterr().out

    monkeypatch.setattr(demo.settings, "DATABASE_URL", "postgresql+psycopg://user@host/db")

    with pytest.raises(SystemExit):
        demo.main()


def test_demo_script_generates_the_xlsx_on_sqlite(db_session, tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(demo, "BACKEND_DIR", tmp_path)

    assert demo.main() == 0

    out = capsys.readouterr().out
    files = list((tmp_path / "var").glob("comparativo-demo-mamposteria-*.xlsx"))
    assert len(files) == 1 and str(files[0]) in out
    assert "[Menor costo total]" in out and "[Menos proveedores]" in out  # L5f
    assert "IVA sin confirmar" in out  # la cal de Ferretería Oeste
    assert load_workbook(files[0]).sheetnames == ["Matriz", "Estrategias", "Supuestos"]
