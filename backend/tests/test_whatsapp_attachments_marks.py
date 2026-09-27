"""Lab L5e (fase 3): marca from_attachment en risk_flags y su paso al comparativo."""

# ruff: noqa: F811 - fixtures importados y recibidos como parámetro

from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace as NS

from openpyxl import load_workbook

from app.features.batch_comparison import fiscal
from app.features.batch_comparison import strategies
from app.features.whatsapp.attachments import marks
from app.features.whatsapp.model import WhatsappMessage
from app.features.whatsapp.quote_writer import evidence_is_literal
from app.features.whatsapp.service import handle_inbound
from tests.test_batch_comparison import _cost
from tests.test_batch_comparison import _rfq
from tests.test_batch_comparison_api import HEADERS
from tests.test_batch_comparison_api import world as comparison_world  # noqa: F401 - fixture
from tests.test_whatsapp_agent import FakeClient
from tests.test_whatsapp_agent import _greet
from tests.test_whatsapp_agent import _quotes
from tests.test_whatsapp_agent import record
from tests.test_whatsapp_agent import reply
from tests.test_whatsapp_agent import text
from tests.test_whatsapp_agent import world  # noqa: F401 - fixture

ATTACHMENT_BODY = "[adjunto: lista.pdf]\nCemento CPN40 x 50 kg $ 12.900 + IVA"


def _attachment_inbound(db, conversation, body=ATTACHMENT_BODY, wa_id="wamid.pdf"):
    message = WhatsappMessage(
        conversation_id=conversation.id, direction="inbound", body=body, wa_message_id=wa_id,
        guardrail_flags=[marks.FLAG_PDF, marks.FLAG_TRANSCRIBED], media_type="document",
    )
    db.add(message)
    db.commit()
    db.refresh(message)

    return message


# ------------------------------------------------------------------ unidad
def test_apply_from_attachment_adds_keeps_and_removes():
    quote = NS(risk_flags=["Complete submission."])
    text_bodies = ["hola, te paso precios"]
    attachment = ["[adjunto: lista.pdf]\nCemento CPN40 $ 12.900 + IVA"]

    assert marks.apply_from_attachment(quote, "Cemento CPN40 $ 12.900", text_bodies + attachment, attachment, literal=evidence_is_literal)
    assert quote.risk_flags == ["Complete submission.", "from_attachment"]

    # Ya está: no se duplica.
    marks.apply_from_attachment(quote, "12.900 + IVA", text_bodies + attachment, attachment, literal=evidence_is_literal)
    assert quote.risk_flags.count("from_attachment") == 1

    # Literal en un texto del proveedor: se quita, sin pisar lo demás.
    assert not marks.apply_from_attachment(quote, "te paso precios", text_bodies + attachment, attachment, literal=evidence_is_literal)
    assert quote.risk_flags == ["Complete submission."]

    # Sin adjuntos ni marca: risk_flags queda como estaba (None).
    empty = NS(risk_flags=None)
    assert not marks.apply_from_attachment(empty, "te paso precios", text_bodies, [], literal=evidence_is_literal)
    assert empty.risk_flags is None


# ------------------------------------------------------------------ agente
def test_price_recorded_from_a_transcribed_attachment_is_marked(db_session, world):
    conversation, _ = _greet(db_session, world)
    attachment = _attachment_inbound(db_session, conversation)
    client = FakeClient([
        reply(record(world.cemento.id, "Cemento CPN40 x 50 kg $ 12.900 + IVA", price=12900, iva_included=False)),
        reply(text("Lo tengo.")),
    ])

    result = handle_inbound(db_session, conversation.id, "", client=client, wa_message_id="wamid.pdf", existing_inbound=attachment)

    assert result.outbound.body == "Lo tengo."
    quote = _quotes(db_session, conversation)[0]
    assert quote.unit_price == Decimal("12900")
    assert quote.risk_flags == ["from_attachment"]


def test_mark_is_removed_when_the_supplier_confirms_in_text(db_session, world):
    conversation, _ = _greet(db_session, world)
    attachment = _attachment_inbound(db_session, conversation)
    handle_inbound(
        db_session, conversation.id, "", wa_message_id="wamid.pdf", existing_inbound=attachment,
        client=FakeClient([reply(record(world.cemento.id, "Cemento CPN40 x 50 kg $ 12.900 + IVA", price=12900)), reply(text("Lo tengo."))]),
    )
    assert _quotes(db_session, conversation)[0].risk_flags == ["from_attachment"]

    handle_inbound(
        db_session, conversation.id, "Sí, cemento 12.900 + IVA, confirmado", wa_message_id="wamid.txt",
        client=FakeClient([reply(record(world.cemento.id, "cemento 12.900 + IVA", price=12900)), reply(text("Perfecto, queda registrado."))]),
    )

    assert _quotes(db_session, conversation)[0].risk_flags == []


def test_price_from_text_only_is_not_marked(db_session, world):
    conversation, _ = _greet(db_session, world)
    _attachment_inbound(db_session, conversation, body="[adjunto: foto.jpg]\nCal 6.200")

    handle_inbound(
        db_session, conversation.id, "Cemento 12.900 con IVA", wa_message_id="wamid.t",
        client=FakeClient([reply(record(world.cemento.id, "Cemento 12.900 con IVA", price=12900)), reply(text("Lo tengo."))]),
    )

    assert _quotes(db_session, conversation)[0].risk_flags is None


# --------------------------------------------------------------- comparativo
def test_fiscal_reads_the_mark_from_risk_flags():
    cost = _cost("1210", risk_flags=["from_attachment"])

    assert cost.from_attachment is True
    assert fiscal.MARK_FROM_ATTACHMENT in cost.marks
    assert _cost("1210").from_attachment is False and fiscal.MARK_FROM_ATTACHMENT not in _cost("1210").marks


def test_porque_caveat_for_prices_read_from_an_attachment():
    items = [_rfq(1, "Cemento", 10), _rfq(2, "Cal", 5)]
    names = {"id:1": "Corralón Norte", "id:2": "Materiales del Sur"}
    costs = [
        _cost("100", items[0], "id:1", iva_included=False, risk_flags=["from_attachment"]),
        _cost("50", items[1], "id:1", iva_included=False, risk_flags=["from_attachment"]),
        _cost("120", items[0], "id:2", iva_included=False),
        _cost("60", items[1], "id:2", iva_included=False),
    ]

    porque = strategies.compare(items, costs, names).lowest_cost.porque

    assert porque.endswith("Corralón Norte: el precio de Cemento y Cal se leyó de un adjunto, revisar contra el archivo.")

    costs[0].marks.insert(0, fiscal.MARK_FREIGHT_TO_QUOTE)
    porque = strategies.compare(items, costs, names).lowest_cost.porque
    assert porque.endswith(
        "Corralón Norte no incluyó el flete en Cemento: el total no lo contempla; además, el precio de Cemento y Cal "
        "se leyó de un adjunto: revisar contra el archivo."
    )


def test_comparison_json_and_excel_show_the_mark(client, db_session, comparison_world):
    quote = next(q for q in comparison_world.ladrillo.quotes if q.supplier_id == comparison_world.sur.id)
    quote.risk_flags = ["from_attachment"]
    db_session.commit()
    sur = f"id:{comparison_world.sur.id}"

    body = client.get(f"/rfq-batches/{comparison_world.batch.id}/comparison", headers=HEADERS).json()
    ladrillo = body["items"][0]
    assert ladrillo["quotes"][sur]["from_attachment"] is True
    assert ladrillo["quotes"][f'id:{comparison_world.norte.id}']["from_attachment"] is False
    assert "precio leído de un adjunto" in ladrillo["quotes"][sur]["marks"]
    assert "se leyó de un adjunto" in body["strategies"][0]["porque"]

    response = client.get(f"/rfq-batches/{comparison_world.batch.id}/comparison.xlsx", headers=HEADERS)
    workbook = load_workbook(BytesIO(response.content))
    matrix_marks = workbook["Matriz"].cell(row=3, column=25).value  # columna Marcas de Materiales del Sur (L5f: +Régimen)
    assert "precio leído de un adjunto" in matrix_marks
    assumptions = [str(c.value) for row in workbook["Supuestos"].iter_rows() for c in row if c.value]
    assert any("Precios leídos de un adjunto (PDF o foto): revisar contra el archivo" in t for t in assumptions)
