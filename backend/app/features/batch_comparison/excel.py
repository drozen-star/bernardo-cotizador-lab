"""Excel del comparativo (L5d) con openpyxl: hojas Matriz, Estrategias y Supuestos.

Los montos se escriben como números (``Decimal`` redondeado a dos decimales) con formato
``#,##0.00``, nunca como texto. Encabezados en castellano. El mejor costo real de cada ítem
se resalta en la matriz.
"""

import re
import unicodedata
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from app.features.batch_comparison import fiscal
from app.features.batch_comparison.service import BUENOS_AIRES
from app.features.batch_comparison.service import BatchComparison
from app.features.batch_comparison.strategies import StrategyResult

MONEY_FORMAT = "#,##0.00"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

BOLD = Font(bold=True)
BEST_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
WRAP = Alignment(wrap_text=True, vertical="top")

ITEM_HEADERS = ("Ítem", "Cantidad", "Unidad")
SUPPLIER_HEADERS = (
    "Precio cotizado",
    "IVA",
    "Neto unitario",
    "Total costo real",
    "Total desembolso",
    "Flete",
    "Plazo (días)",
    "Pago",
    "Validez",
    "Marcas",
)

INVOICED_NOTE = (
    "Todas las cotizaciones se asumen facturadas (con factura A). "
    "Si alguna es en efectivo, este comparativo no la compara bien."
)
BETA_NOTE = (
    "Desembolso = neto × (1 + alícuota): la plata que sale. "
    "Costo real (β=1) = neto: el IVA se computa como crédito fiscal contra el débito de la obra. "
    "Si no recuperás el IVA (β=0), mirá el desembolso."
)
IVA_UNCONFIRMED_NOTE = "Si el proveedor no aclaró si el precio incluye IVA, se toma sin IVA (el caso más caro) y se marca."
FREIGHT_NOTE = (
    "Si el proveedor pasó un costo de flete, se suma al neto del ítem. "
    "'Flete a cotizar' significa que el precio no lo incluye y no hay costo: el total queda sin flete."
)
CURRENCY_NOTE = "Moneda: ARS; otras monedas no se comparan."


# ------------------------------------------------------------------ helpers
def _money_cell(ws, row: int, column: int, value: Decimal | None) -> None:
    if value is None:
        return

    cell = ws.cell(row=row, column=column, value=fiscal.money(value))
    cell.number_format = MONEY_FORMAT


def _freight_text(cost) -> Decimal | str:
    if cost.shipping_cost is not None:
        return fiscal.money(cost.shipping_cost)

    return {True: "incluido", False: "a cotizar", None: "sin confirmar"}[cost.freight_included]


def slugify(text: str, max_length: int = 40) -> str:
    ascii_text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_text.lower()).strip("-")

    return slug[:max_length].strip("-") or "pedido"


def filename_for(batch_name: str, when: datetime) -> str:
    return f"comparativo-{slugify(batch_name)}-{when.astimezone(BUENOS_AIRES).date().isoformat()}.xlsx"


# --------------------------------------------------------------------- hojas
def _matrix_sheet(ws, result: BatchComparison) -> None:
    loaded = result.loaded
    width = len(SUPPLIER_HEADERS)

    for index, header in enumerate(ITEM_HEADERS, start=1):
        ws.cell(row=2, column=index, value=header).font = BOLD

    for s_index, supplier in enumerate(loaded.suppliers):
        first = len(ITEM_HEADERS) + 1 + s_index * width
        title = supplier.name

        if supplier.conversation_status:
            title += f" (WhatsApp: {supplier.conversation_status})"

        ws.cell(row=1, column=first, value=title).font = BOLD
        ws.merge_cells(start_row=1, start_column=first, end_row=1, end_column=first + width - 1)

        for offset, header in enumerate(SUPPLIER_HEADERS):
            ws.cell(row=2, column=first + offset, value=header).font = BOLD

    for r_index, rfq in enumerate(loaded.items, start=3):
        ws.cell(row=r_index, column=1, value=rfq.item_name)
        ws.cell(row=r_index, column=2, value=rfq.quantity)
        ws.cell(row=r_index, column=3, value=rfq.unit)
        best_key = result.best_supplier_key(rfq.id)

        for s_index, supplier in enumerate(loaded.suppliers):
            cost = result.cost_for(rfq.id, supplier.key)

            if cost is None:
                continue

            first = len(ITEM_HEADERS) + 1 + s_index * width
            _money_cell(ws, r_index, first, cost.unit_price)
            ws.cell(row=r_index, column=first + 1, value=fiscal.IVA_LABELS[cost.iva_included])
            _money_cell(ws, r_index, first + 2, cost.neto_unit)
            _money_cell(ws, r_index, first + 3, cost.costo_real_total)
            _money_cell(ws, r_index, first + 4, cost.desembolso_total)
            freight = _freight_text(cost)
            freight_cell = ws.cell(row=r_index, column=first + 5, value=freight)

            if isinstance(freight, Decimal):
                freight_cell.number_format = MONEY_FORMAT

            ws.cell(row=r_index, column=first + 6, value=cost.lead_time)
            ws.cell(row=r_index, column=first + 7, value=cost.payment_terms)
            ws.cell(row=r_index, column=first + 8, value=cost.validity)
            ws.cell(row=r_index, column=first + 9, value="; ".join(cost.marks) or None)

            if supplier.key == best_key and cost.comparable:
                ws.cell(row=r_index, column=first + 3).fill = BEST_FILL

    ws.column_dimensions["A"].width = 38

    for column in range(4, 4 + width * len(loaded.suppliers)):
        ws.column_dimensions[get_column_letter(column)].width = 16


def _strategy_block(ws, row: int, result: StrategyResult) -> int:
    ws.cell(row=row, column=1, value=result.title).font = Font(bold=True, size=12)
    row += 1

    for index, header in enumerate(("Ítem", "Proveedor", "Costo real", "Desembolso", "Plazo (días)", "Marcas"), 1):
        ws.cell(row=row, column=index, value=header).font = BOLD

    row += 1

    for assignment in result.assignments:
        ws.cell(row=row, column=1, value=assignment.item_name)
        ws.cell(row=row, column=2, value=assignment.supplier_name or "sin cotización")

        if assignment.cost is not None:
            _money_cell(ws, row, 3, assignment.cost.costo_real_total)
            _money_cell(ws, row, 4, assignment.cost.desembolso_total)
            ws.cell(row=row, column=5, value=assignment.cost.lead_time)
            ws.cell(row=row, column=6, value="; ".join(assignment.cost.marks) or None)

        row += 1

    ws.cell(row=row, column=1, value="Total costo real").font = BOLD
    _money_cell(ws, row, 3, result.total_costo_real)
    ws.cell(row=row + 1, column=1, value="Total desembolso").font = BOLD
    _money_cell(ws, row + 1, 4, result.total_desembolso)
    ws.cell(row=row + 2, column=1, value="Cantidad de proveedores").font = BOLD
    ws.cell(row=row + 2, column=2, value=result.supplier_count)
    ws.cell(row=row + 3, column=1, value="Plazo máximo (días)").font = BOLD
    ws.cell(row=row + 3, column=2, value=result.max_lead_time)
    ws.cell(row=row + 4, column=1, value="Marcas").font = BOLD
    ws.cell(row=row + 4, column=2, value="; ".join(result.marks) or "ninguna").alignment = WRAP
    ws.cell(row=row + 5, column=1, value="Por qué").font = BOLD
    ws.cell(row=row + 5, column=2, value=result.porque).alignment = WRAP

    return row + 7


def _strategies_sheet(ws, result: BatchComparison) -> None:
    row = 1

    for strategy in result.comparison.strategies:
        row = _strategy_block(ws, row, strategy)

    comparison = result.comparison
    ws.cell(row=row, column=1, value="Diferencia entre estrategias").font = Font(bold=True, size=12)
    ws.cell(row=row + 1, column=1, value="Menos proveedores paga de más (costo real)")
    _money_cell(ws, row + 1, 3, comparison.difference_amount)
    ws.cell(row=row + 2, column=1, value="Diferencia porcentual")
    ws.cell(
        row=row + 2, column=2,
        value=None if comparison.difference_pct is None else fiscal.format_pct(comparison.difference_pct),
    )

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 60

    for letter in ("C", "D", "E", "F"):
        ws.column_dimensions[letter].width = 18


def _assumptions_sheet(ws, result: BatchComparison) -> None:
    batch = result.loaded.batch
    site = ", ".join(part for part in (batch.site_name, batch.site_address) if part) or "sin obra informada"
    generated = result.generated_at.astimezone(BUENOS_AIRES).strftime("%d/%m/%Y %H:%M") + " (Buenos Aires)"

    rows: list[tuple[str, str]] = [
        ("Pedido", batch.name),
        ("Obra", site),
        ("Generado", generated),
        ("Facturación", INVOICED_NOTE),
        ("Desembolso vs costo real", BETA_NOTE),
        ("IVA sin confirmar", IVA_UNCONFIRMED_NOTE),
        ("Flete", FREIGHT_NOTE),
        ("Moneda", CURRENCY_NOTE),
    ]

    for rfq in result.loaded.items:
        rows.append((f"Alícuota IVA: {rfq.item_name}", f"{result.alicuotas[rfq.id]}%"))

    for index, (key, value) in enumerate(rows, start=1):
        ws.cell(row=index, column=1, value=key).font = BOLD
        ws.cell(row=index, column=2, value=value).alignment = WRAP

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 100


# ------------------------------------------------------------------ público
def build_workbook(result: BatchComparison) -> Workbook:
    workbook = Workbook()
    matrix = workbook.active
    matrix.title = "Matriz"
    _matrix_sheet(matrix, result)
    _strategies_sheet(workbook.create_sheet("Estrategias"), result)
    _assumptions_sheet(workbook.create_sheet("Supuestos"), result)

    return workbook


def workbook_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)

    return buffer.getvalue()
