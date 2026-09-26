"""Intake de materiales desde un Excel de formato fijo (laboratorio Bernardo, L2).

Lectura determinística con openpyxl: sin modelo de lenguaje ni heurísticas. El formato
es un contrato:

* hoja 1 (la primera del libro),
* fila 1 con los encabezados exactos ``item, quantity, unit, specification,
  accepted_alternatives`` (en cualquier orden, sin columnas extra),
* una fila por material, de la 2 en adelante.

Todo desvío se reporta con el número de fila de Excel, y se reportan **todos** los
problemas en una sola pasada, para que el comprador corrija el archivo una vez y no
un error por intento.

Decisión sobre ``quantity``: además de numérica y > 0 tiene que ser **entera**, porque
``rfqs.quantity`` en la base es ``Integer``. Una cantidad como 2,5 m³ se expresa
cambiando la unidad (2500 litros) o redondeando en el Excel; acá no se redondea nada
en silencio.
"""

import io
import zipfile
from decimal import Decimal
from decimal import InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.exceptions import BadRequestError

EXPECTED_HEADERS: tuple[str, ...] = (
    "item",
    "quantity",
    "unit",
    "specification",
    "accepted_alternatives",
)

#: Filas de datos como máximo. Un pedido de obra real anda en decenas de ítems.
MAX_ROWS = 500


class MaterialsIntakeError(BadRequestError):
    """El Excel tiene uno o más problemas. ``errors`` trae uno por línea, con fila."""

    def __init__(self, errors: list[str]) -> None:
        self.errors = list(errors)
        super().__init__("El Excel tiene errores: " + " | ".join(self.errors))


def _cell_text(value: Any) -> str:
    if value is None:
        return ""

    return str(value).strip()


def _parse_quantity(value: Any) -> Decimal | None:
    """Número o ``None`` si no se puede leer como número.

    Acepta int, float, Decimal y texto con punto o coma decimal ("1.200" NO se
    interpreta como mil doscientos: el punto es decimal, como lo guarda Excel).
    """

    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))

    text = str(value).strip().replace(",", ".")

    if not text:
        return None

    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _open_first_sheet(file: Any):
    if isinstance(file, (bytes, bytearray)):
        source: Any = io.BytesIO(file)
    elif isinstance(file, (str, Path)):
        source = str(file)
    else:
        source = file  # file-like abierto en binario

    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise MaterialsIntakeError(["El archivo no es un .xlsx válido."]) from exc

    if not workbook.worksheets:
        workbook.close()
        raise MaterialsIntakeError(["El libro no tiene hojas."])

    return workbook, workbook.worksheets[0]


def _read_headers(header_row: tuple | None) -> tuple[dict[str, int], list[str]]:
    """Mapa encabezado -> índice de columna, más los errores de encabezado."""

    errors: list[str] = []
    cells = [_cell_text(cell).lower() for cell in (header_row or ())]

    columns: dict[str, int] = {}

    for index, header in enumerate(cells):
        if not header:
            continue

        if header not in EXPECTED_HEADERS:
            errors.append(f"fila 1: columna desconocida '{header}'")
        elif header in columns:
            errors.append(f"fila 1: encabezado repetido '{header}'")
        else:
            columns[header] = index

    missing = [header for header in EXPECTED_HEADERS if header not in columns]

    if missing:
        errors.append(
            "fila 1: faltan los encabezados "
            + ", ".join(f"'{header}'" for header in missing)
            + ". Se esperan exactamente: "
            + ", ".join(EXPECTED_HEADERS)
        )

    return columns, errors


def _is_blank(row: tuple) -> bool:
    return all(_cell_text(cell) == "" for cell in row)


def parse_materials_xlsx(file: Any) -> list[dict]:
    """Lee el Excel y devuelve una lista de dicts, uno por material.

    ``file`` puede ser bytes, una ruta o un objeto binario abierto. Cada dict trae
    ``row`` (fila de Excel), ``item``, ``quantity`` (int), ``unit``,
    ``specification`` (str, puede ser vacío) y ``accepted_alternatives``
    (str o ``None``). Lanza :class:`MaterialsIntakeError` con todos los problemas.
    """

    workbook, sheet = _open_first_sheet(file)

    try:
        rows_iter = sheet.iter_rows(values_only=True)
        header_row = next(rows_iter, None)

        columns, errors = _read_headers(header_row)

        parsed: list[dict] = []
        seen: dict[str, int] = {}
        data_rows = 0

        def cell(row: tuple, name: str) -> Any:
            index = columns.get(name)

            if index is None or index >= len(row):
                return None

            return row[index]

        for row_number, row in enumerate(rows_iter, start=2):
            if _is_blank(row):
                continue

            data_rows += 1

            if data_rows > MAX_ROWS:
                errors.append(f"fila {row_number}: el Excel supera las {MAX_ROWS} filas")
                break

            item = _cell_text(cell(row, "item"))
            unit = _cell_text(cell(row, "unit"))
            specification = _cell_text(cell(row, "specification"))
            alternatives = _cell_text(cell(row, "accepted_alternatives")) or None

            if "item" in columns and not item:
                errors.append(f"fila {row_number}: item vacío")

            if "unit" in columns and not unit:
                errors.append(f"fila {row_number}: unit vacío")

            quantity: int | None = None

            if "quantity" in columns:
                raw = cell(row, "quantity")
                number = _parse_quantity(raw)

                if number is None:
                    errors.append(
                        f"fila {row_number}: quantity debe ser numérica (recibido '{_cell_text(raw)}')"
                    )
                elif number <= 0:
                    errors.append(
                        f"fila {row_number}: quantity debe ser mayor que 0 (recibido {raw})"
                    )
                elif number != number.to_integral_value():
                    errors.append(
                        f"fila {row_number}: quantity debe ser entera (recibido {raw}); "
                        "la base guarda cantidades enteras, ajustá la unidad"
                    )
                else:
                    quantity = int(number)

            if item:
                key = " ".join(item.lower().split())

                if key in seen:
                    errors.append(
                        f"fila {row_number}: ítem duplicado '{item}' "
                        f"(ya aparece en la fila {seen[key]})"
                    )
                else:
                    seen[key] = row_number

            parsed.append(
                {
                    "row": row_number,
                    "item": item,
                    "quantity": quantity,
                    "unit": unit,
                    "specification": specification,
                    "accepted_alternatives": alternatives,
                }
            )

        if data_rows == 0:
            errors.append("El Excel no tiene filas de materiales debajo del encabezado.")

    finally:
        workbook.close()

    if errors:
        raise MaterialsIntakeError(errors)

    return parsed
